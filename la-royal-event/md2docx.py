#!/usr/bin/env python3
"""Markdown -> DOCX на python-docx. Заголовки, абзацы, списки, таблицы, цитаты, выделения."""
import sys, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VIOLET = RGBColor(0x4A, 0x2D, 0x7A)
GREY = RGBColor(0x55, 0x55, 0x55)

INLINE = re.compile(r'(\*\*.+?\*\*|__.+?__|\*[^*]+?\*|_[^_]+?_|`[^`]+?`)')


def setup(doc):
    st = doc.styles['Normal']
    st.font.name = 'Georgia'
    st.font.size = Pt(11)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Georgia')
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.4)
        s.right_margin = Cm(2)
    for name, size, color, bold in (('Heading 1', 20, VIOLET, True),
                                    ('Heading 2', 14.5, VIOLET, True),
                                    ('Heading 3', 12.5, None, True),
                                    ('Heading 4', 11, None, True),
                                    ('Heading 5', 10.5, GREY, True)):
        try:
            h = doc.styles[name]
        except KeyError:
            continue
        h.font.name = 'Georgia'
        h.font.size = Pt(size)
        h.font.bold = bold
        if color is not None:
            h.font.color.rgb = color
        h.paragraph_format.space_before = Pt(14 if size > 13 else 10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True


def add_runs(par, text):
    """Разбирает **жирный**, *курсив* и `код`."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if (chunk.startswith('**') and chunk.endswith('**')) or (chunk.startswith('__') and chunk.endswith('__')):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith('`') and chunk.endswith('`'):
            r = par.add_run(chunk[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
        elif (chunk.startswith('*') and chunk.endswith('*') and len(chunk) > 2) or \
             (chunk.startswith('_') and chunk.endswith('_') and len(chunk) > 2):
            par.add_run(chunk[1:-1]).italic = True
        else:
            par.add_run(chunk)


def add_table(doc, rows):
    header, body = rows[0], rows[2:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, cell in enumerate(header):
        p = t.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, cell)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
        shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'EFE9F5')
        t.rows[0].cells[i]._tc.get_or_add_tcPr().append(shd)
    for row in body:
        cells = t.add_row().cells
        for i, cell in enumerate(row[:len(header)]):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, cell)
            for r in p.runs:
                r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


def convert(md_path, out_path=None, title=None):
    text = open(md_path, encoding='utf-8').read()
    out_path = out_path or os.path.splitext(md_path)[0] + '.docx'
    doc = Document()
    setup(doc)
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # таблица
        if line.lstrip().startswith('|') and i + 1 < len(lines) and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        # заголовок
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            lvl = min(len(m.group(1)), 5)
            p = doc.add_paragraph(style=f'Heading {lvl}')
            add_runs(p, m.group(2))
            i += 1
            continue
        # горизонтальная линия
        if re.match(r'^\s*(---+|\*\*\*+|___+)\s*$', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pbd = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
            bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '4')
            bot.set(qn('w:space'), '1'); bot.set(qn('w:color'), 'C8BFD4')
            pbd.append(bot); p._p.get_or_add_pPr().append(pbd)
            i += 1
            continue
        # цитата
        if line.lstrip().startswith('> '):
            buf = []
            while i < len(lines) and lines[i].lstrip().startswith('>'):
                buf.append(lines[i].lstrip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_runs(p, ' '.join(buf))
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = GREY
            continue
        # маркированный / нумерованный список
        m = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.*)$', line)
        if m:
            indent = len(m.group(1)) // 2
            ordered = m.group(2)[0].isdigit()
            style = 'List Number' if ordered else 'List Bullet'
            p = doc.add_paragraph(style=style)
            if indent:
                p.paragraph_format.left_indent = Cm(0.75 + 0.5 * indent)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, m.group(3))
            i += 1
            continue
        # обычный абзац
        buf = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r'^(#{1,6}\s|\s*[-*+]\s|\s*\d+\.\s|\s*\||>|\s*---+\s*$)', lines[i]):
            buf.append(lines[i].rstrip())
            i += 1
        p = doc.add_paragraph()
        add_runs(p, ' '.join(buf))

    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    for p in sys.argv[1:]:
        o = convert(p)
        print(f'{os.path.basename(p)} -> {os.path.basename(o)} ({os.path.getsize(o)//1024} КБ)')
